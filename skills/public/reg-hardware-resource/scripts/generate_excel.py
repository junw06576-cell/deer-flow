#!/usr/bin/env python3
"""
WiNEX 区域卫生硬件资源清单生成器 v2.0

用法:
  python3 generate_excel.py --config config.json --source-dir <平铺md目录> [--output output.xlsx]

--source-dir: 存放平铺格式源数据 md 的目录（由 xlsx_to_md.py 从 Excel 转换而来，
即知识库 ingest 流程生成的笔记 md）。目录下所有 .md 均会被读取，
按 `## {sheet名}` 定位产品数据，sheet 名需含（信创）/（非信创）标识。

config.json 格式:
{
  "project_name": "项目名称",
  "scope": "建设范围",
  "population": 100,
  "xinchuang": true,
  "products": {
    "基层医疗": {"outpatient": 8000},
    "智慧公卫": {"variant": "基本公卫+签约"},
    "区域平台": {"combo": "全量清单"}
  }
}

products 中可用的产品键:
  - 基层医疗: 需要 outpatient 参数 (日门诊量, 整数)
  - 智慧公卫: 需要 variant 参数 (基本公卫|签约|健康管理|基本公卫+签约|基本公卫+签约+健康管理)
  - 健康体检: 无额外参数
  - 区域平台: 需要 combo 参数 (全量清单|基础EHR+BI|检查检验互认|双向转诊|健康档案共享|综合监管)

源数据目录中的 md 平铺格式要求（xlsx_to_md.py 自动保证）:
  - 每个 Excel sheet 对应一个 `## {sheet名}` 二级标题
  - sheet 名: WXP必选 | 基层医疗（信创/非信创）| 智慧公卫（…）| 健康体检（…）| 区域平台（…）
  - sheet 内每个档位/变体一个章节标题行（原 Excel 横向合并单元格，展开后整行同值）
  - 数据行 17 列，首列为 生产环境/测试环境
"""

import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
    from openpyxl.utils import get_column_letter
except ImportError:
    print("错误: 需要安装 openpyxl。运行: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

SCRIPT_DIR = Path(__file__).resolve().parent

COL_HEADERS = [
    "部署类型", "产品", "服务明细", "选项", "资源类型",
    "服务器数量", "CPU（核）", "内存（GB）", "系统盘（GB）", "数据盘（GB）",
    "磁盘要求", "操作系统版本", "应用中间件", "数据库", "开放端口", "建议带宽", "备注"
]

OUTPUT_COL_HEADERS = [
    "部署类型", "产品", "服务明细", "选项", "资源类型",
    "服务器数量", "CPU（核）", "内存（GB）", "系统盘（GB）", "数据盘（GB）",
    "磁盘要求", "操作系统版本", "应用中间件", "数据库", "开放端口", "备注"
]

NUM_OUTPUT_COLS = len(OUTPUT_COL_HEADERS)

NUM_FMT = '#,##0'

THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
HEADER_FILL = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')
HEADER_FONT = Font(bold=True)
TITLE_FONT = Font(bold=True, size=14)
SUBTITLE_FONT = Font(bold=True, size=11)
SUMMARY_FILL = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')
SUMMARY_FONT = Font(bold=True, size=11)

PRODUCT_SHEET_NAMES = {
    "基层医疗": "基层医疗信息系统",
    "智慧公卫": "智慧公卫系统",
    "健康体检": "健康体检",
}

REGIONAL_SHEET_NAMES = {
    "全量清单": "区域平台",
    "基础EHR+BI": "区域平台-基础EHR+BI",
    "检查检验互认": "区域平台-检查检验互认",
    "双向转诊": "区域平台-双向转诊",
    "健康档案共享": "区域平台-健康档案共享",
    "综合监管": "区域平台-综合监管",
}

# 数据源 sheet 名模板（{xc} = 信创/非信创）
SHEET_NAME_TEMPLATES = {
    "基层医疗": "基层医疗（{xc}）",
    "智慧公卫": "智慧公卫（{xc}）",
    "健康体检": "健康体检（{xc}）",
    "区域平台": "区域平台（{xc}）",
    "WXP": "WXP必选",
}

# 平铺源数据全局缓存: {sheet名: [(章节标题|None, [行dict])]}
FLAT_SHEETS = {}

# 适配软件版本表的产品列值（这些行不是资源配置数据）
ADAPTER_PRODS = {'分类', '数据库', '服务器', '客户端', '浏览器', '中间件'}

DATA_ROW_PREFIXES = ('生产环境', '测试环境', '正式环境')

TIER_KEYWORDS = {
    "基层医疗": {
        0: "非标最小",
        1: "1000<日门诊量≤5000",
        2: "5000<日门诊量≤10000",
        3: "10000<日门诊量≤15000",
    },
    "公卫类": {
        1: "区域人口数≤100万",
        2: "100万<区域人口数≤300万",
        3: "300万<区域人口数≤700万",
    },
    "区域平台": {
        1: "人口100万以下",
        2: "人口100万~300万",
        3: "人口300万以上",
    },
}

PUBLIC_HEALTH_PRODUCTS = {"健康体检"}
SHARED_COMPONENT_PRODUCTS = {"智慧公卫", "健康体检"}

COL_WIDTHS = {
    'A': 12, 'B': 25, 'C': 35, 'D': 18, 'E': 14,
    'F': 12, 'G': 10, 'H': 10, 'I': 12, 'J': 12,
    'K': 10, 'L': 30, 'M': 25, 'N': 20, 'O': 16, 'P': 35
}


def strip_version(text):
    """去除产品名中的版本号，如 V5.6、V6.0、 5.6"""
    if not text:
        return text
    text = re.sub(r'\s*V?\d+\.\d+\s*', '', str(text))
    return re.sub(r'\s+', ' ', text).strip()


def determine_tier(product_key, outpatient=None, population=None, minimal=False):
    if product_key == "基层医疗":
        if outpatient is None:
            raise ValueError("基层医疗需要 outpatient 参数")
        if minimal and outpatient <= 1000:
            return 0
        if outpatient <= 5000:
            return 1
        elif outpatient <= 10000:
            return 2
        elif outpatient <= 15000:
            return 3
        else:
            raise ValueError(f"日门诊量 {outpatient} 超出范围 (最大15000)")
    elif product_key == "区域平台":
        if population is None:
            raise ValueError("区域平台需要 population 参数")
        if population <= 100:
            return 1
        elif population <= 300:
            return 2
        else:
            return 3
    elif product_key == "智慧公卫":
        if population is None:
            raise ValueError("智慧公卫需要 population 参数")
        if population <= 100:
            return 1
        elif population <= 300:
            return 2
        else:
            return 3
    elif product_key in PUBLIC_HEALTH_PRODUCTS:
        if population is None:
            raise ValueError(f"{product_key}需要 population 参数")
        if population <= 100:
            return 1
        elif population <= 300:
            return 2
        else:
            return 3
    else:
        raise ValueError(f"未知产品: {product_key}")


TIER_FALLBACK_NOTES = {}


def load_flat_source(source):
    """加载数据源：本地目录（读取全部 *.md）或 HTTP(S) URL（单文件，如知识库数据源笔记）"""
    src = str(source)
    FLAT_SHEETS.clear()
    if src.startswith(('http://', 'https://')):
        import urllib.request
        from urllib.parse import quote
        url = quote(src, safe=':/?&=%#')
        with urllib.request.urlopen(url, timeout=60) as resp:
            text = resp.read().decode('utf-8')
        if not text.strip():
            raise ValueError(f"数据源 URL 返回空内容: {src}")
        for sheet, sections in parse_flat_md_text(text).items():
            FLAT_SHEETS.setdefault(sheet, []).extend(sections)
        return
    path = Path(src)
    if not path.is_dir():
        raise FileNotFoundError(f"源数据目录不存在: {path}")
    md_files = sorted(path.glob('*.md'))
    if not md_files:
        raise FileNotFoundError(f"源数据目录中没有 .md 文件: {path}")
    for md in md_files:
        for sheet, sections in parse_flat_md(md).items():
            FLAT_SHEETS.setdefault(sheet, []).extend(sections)


def parse_flat_md(filepath):
    with open(filepath, 'r', encoding='utf-8') as f:
        return parse_flat_md_text(f.read())


def parse_flat_md_text(text):
    """解析 xlsx_to_md.py 产出的平铺 md。

    格式特征：
      - 每个 sheet 一个 `## {sheet名}` 标题，正文是整 sheet 一张 GFM 大表
      - 章节标题行 = 原 Excel 横向合并单元格，展开后整行非空单元格同值
      - 数据行首列为 生产环境/测试环境/正式环境
    返回 {sheet名: [(章节标题|None, [行dict])]}
    """
    sections_by_sheet = {}
    state = {'sheet': None, 'title': None, 'rows': None}

    def flush():
        if state['sheet'] is not None and state['rows']:
            sections_by_sheet.setdefault(state['sheet'], []).append((state['title'], state['rows']))

    for raw in text.split('\n'):
            line = raw.rstrip('\n')
            if line.startswith('## '):
                flush()
                state['sheet'] = line[3:].strip()
                state['title'], state['rows'] = None, None
                continue
            if not line.startswith('|') or state['sheet'] is None:
                continue
            cells = [c.strip() for c in line.strip('|').split('|')]
            if not cells or all(c == '' for c in cells):
                continue
            if all(set(c) <= {'-', ':', ' '} for c in cells if c != ''):
                continue  # 表格分隔行 |---|---|
            first = cells[0]
            nonempty = [c for c in cells if c != '']
            # 章节标题行：整行非空单元格同值（横向合并展开），排除注释/适配表标题
            if (len(set(nonempty)) == 1 and first
                    and first != '部署类型'
                    and not first.startswith('注')
                    and not re.match(r'^\d+、', first)
                    and first != '信创适配软件版本要求'):
                flush()
                state['title'], state['rows'] = first, []
                continue
            if first == '部署类型':
                if state['rows'] is None:
                    state['rows'] = []  # 无章节标题的 sheet（如 WXP必选）
                continue
            if first in DATA_ROW_PREFIXES:
                prod = cells[1] if len(cells) > 1 else ''
                if prod in ADAPTER_PRODS:
                    continue  # 信创适配软件版本表
                if state['rows'] is None:
                    state['rows'] = []
                cells = [re.sub(r'\s+', ' ', c) for c in cells]
                state['rows'].append(_parse_row((cells + [''] * 17)[:17]))
    flush()
    return sections_by_sheet


def _get_tier_keywords(tier):
    all_kw = []
    for cat in TIER_KEYWORDS.values():
        if tier in cat:
            all_kw.append(cat[tier])
    all_kw.append(f"第{tier}档")
    return all_kw


def select_rows(sheet_name, tier, variant=None, combo=None):
    """在指定 sheet 的章节中按档位(+变体/套组)选取数据行，支持回退低档"""
    if tier == 0:
        raise ValueError("数据源不含「非标最小档」章节；如需最小档请在源 Excel 中补充该档位")
    sections = FLAT_SHEETS.get(sheet_name)
    if not sections:
        raise FileNotFoundError(
            f"数据源中不存在 sheet「{sheet_name}」，请检查 --source-dir 及源 Excel 的 sheet 名")
    for t in range(tier, 0, -1):
        kws = _get_tier_keywords(t)
        for title, rows in sections:
            if title is None:
                continue
            tt = title.replace('～', '~')
            if not any(kw in tt for kw in kws):
                continue
            if variant and f"（{variant}）" not in tt:
                continue
            if combo and combo not in tt:
                continue
            if t != tier:
                print(f"警告: 档位{tier}未在「{sheet_name}」中找到，回退使用档位{t}", file=sys.stderr)
                TIER_FALLBACK_NOTES[sheet_name] = (tier, t)
            return rows
    raise ValueError(
        f"数据源 sheet「{sheet_name}」中找不到匹配数据 "
        f"(档位{tier}{', 变体=' + variant if variant else ''}{', 套组=' + combo if combo else ''})")


def _parse_row(cells):
    num_cols = 17

    if len(cells) < num_cols:
        cells = cells + [''] * (num_cols - len(cells))
    elif len(cells) > num_cols:
        cells = cells[:num_cols]

    row = {}
    for i, header in enumerate(COL_HEADERS):
        row[header] = cells[i] if i < len(cells) else ''

    for key in ['服务器数量', 'CPU（核）', '内存（GB）', '系统盘（GB）', '数据盘（GB）']:
        val = row.get(key, '')
        if val == '' or val.strip() == '':
            row[key] = None
        else:
            try:
                row[key] = int(val)
            except (ValueError, TypeError):
                try:
                    row[key] = float(val)
                except (ValueError, TypeError):
                    row[key] = None

    if row.get('服务明细', '') == '硬件负载':
        for key in ['服务器数量', 'CPU（核）', '内存（GB）', '系统盘（GB）', '数据盘（GB）']:
            row[key] = None

    return row


def get_tier_description(product_key, tier, outpatient=None, population=None):
    if product_key == "基层医疗":
        tier_map = {
            0: f"非标最小（日门诊量≤1000且机构数量≤3）",
            1: f"第1档（日门诊量≤5000）",
            2: f"第2档（5000<日门诊量≤10000）",
            3: f"第3档（10000<日门诊量≤15000）",
        }
        return tier_map[tier]
    elif product_key == "区域平台":
        tier_map = {
            1: f"第1档（人口100万以下）",
            2: f"第2档（人口100万~300万）",
            3: f"第3档（人口300万以上）",
        }
        pop_str = f" | 区域人口：{population}万" if population else ""
        return tier_map[tier] + pop_str
    elif product_key == "智慧公卫":
        tier_map = {
            1: f"第1档（区域人口≤100万）",
            2: f"第2档（100万<区域人口≤300万）",
            3: f"第3档（区域人口>300万）",
        }
        pop_str = f" | 区域人口：{population}万" if population else ""
        return tier_map[tier] + pop_str
    elif product_key in PUBLIC_HEALTH_PRODUCTS:
        tier_map = {
            1: f"第1档（区域人口≤100万）",
            2: f"第2档（100万<区域人口≤300万）",
            3: f"第3档（区域人口>300万）",
        }
        pop_str = f" | 区域人口：{population}万" if population else ""
        return tier_map[tier] + pop_str
    return ""


def is_shared_component(row):
    product = row.get('产品', '')
    return product in ('公卫公共组件服务器', '公用负载均衡器', '公用负载均衡分发服务器')


def load_product_data(product_key, xinchuang, tier, variant=None):
    sheet_name = SHEET_NAME_TEMPLATES[product_key].format(xc="信创" if xinchuang else "非信创")
    return select_rows(sheet_name, tier, variant=variant)


def load_regional_data(combo, xinchuang, tier):
    sheet_name = SHEET_NAME_TEMPLATES["区域平台"].format(xc="信创" if xinchuang else "非信创")
    return select_rows(sheet_name, tier, combo=combo)


def load_wxp_rows():
    sections = FLAT_SHEETS.get(SHEET_NAME_TEMPLATES["WXP"])
    if not sections or not sections[0][1]:
        raise FileNotFoundError("数据源中不存在 sheet「WXP必选」或该 sheet 无数据")
    return [dict(r, **{'数据库': r.get('数据库') or '-'}) for r in sections[0][1]]


def apply_styles(ws, data_start_row, data_end_row, summary_row):
    for col_idx in range(1, NUM_OUTPUT_COLS + 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = COL_WIDTHS.get(col_letter, 15)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NUM_OUTPUT_COLS)
    title_cell = ws.cell(row=1, column=1)
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NUM_OUTPUT_COLS)
    subtitle_cell = ws.cell(row=2, column=1)
    subtitle_cell.font = SUBTITLE_FONT
    subtitle_cell.alignment = Alignment(horizontal='center', vertical='center')

    for col_idx in range(1, NUM_OUTPUT_COLS + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = THIN_BORDER

    for row_idx in range(data_start_row, data_end_row + 1):
        for col_idx in range(1, NUM_OUTPUT_COLS + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = THIN_BORDER
            cell.font = Font(size=10)
            if col_idx in (6, 7, 8, 9, 10):
                cell.alignment = Alignment(horizontal='center')
            else:
                cell.alignment = Alignment(vertical='center', wrap_text=True)

    for col_idx in range(1, NUM_OUTPUT_COLS + 1):
        cell = ws.cell(row=summary_row, column=col_idx)
        cell.font = SUMMARY_FONT
        cell.fill = SUMMARY_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=5)
    ws.cell(row=summary_row, column=1).value = "合计"
    ws.cell(row=summary_row, column=1).alignment = Alignment(horizontal='center', vertical='center')


def merge_consecutive_cells(ws, col_idx, data_start_row, data_end_row):
    rows = []
    for r in range(data_start_row, data_end_row + 1):
        val = ws.cell(row=r, column=col_idx).value
        rows.append((r, val))

    if not rows:
        return

    merge_start = rows[0][0]
    prev_val = rows[0][1]

    for i in range(1, len(rows)):
        curr_val = rows[i][1]
        if curr_val != prev_val:
            if i - (merge_start - data_start_row) > 1 and merge_start < rows[i - 1][0]:
                ws.merge_cells(
                    start_row=merge_start, start_column=col_idx,
                    end_row=rows[i - 1][0], end_column=col_idx
                )
            merge_start = rows[i][0]
            prev_val = curr_val

    last_idx = len(rows) - 1
    if last_idx - (merge_start - data_start_row) >= 1:
        ws.merge_cells(
            start_row=merge_start, start_column=col_idx,
            end_row=rows[last_idx][0], end_column=col_idx
        )


def write_product_sheet(wb, sheet_name, title, subtitle, rows_data, is_wxp=False):
    ws = wb.create_sheet(title=sheet_name)

    ws.cell(row=1, column=1, value=title)
    ws.cell(row=2, column=1, value=subtitle)

    for col_idx, header in enumerate(OUTPUT_COL_HEADERS, 1):
        ws.cell(row=3, column=col_idx, value=header)

    data_start_row = 4

    for row_offset, row_dict in enumerate(rows_data):
        row_idx = data_start_row + row_offset
        bw = row_dict.get('建议带宽', '') or ''
        remark = row_dict.get('备注', '') or ''
        for col_idx, header in enumerate(OUTPUT_COL_HEADERS, 1):
            val = row_dict.get(header, '')
            if header == '备注' and bw and bw != '-':
                bw_part = f"建议带宽：{bw}"
                val = f"{bw_part}；{remark}" if remark else bw_part
            if val is None:
                val = ''
            if col_idx == 2:
                val = strip_version(val)
            cell = ws.cell(row=row_idx, column=col_idx, value=val if val != '' else None)
            if col_idx in (6, 7, 8, 9, 10) and isinstance(val, (int, float)):
                cell.number_format = NUM_FMT

    data_end_row = data_start_row + len(rows_data) - 1
    summary_row = data_end_row + 1

    ws.cell(row=summary_row, column=6,
            value=f'=SUM(F{data_start_row}:F{data_end_row})')
    ws.cell(row=summary_row, column=6).number_format = NUM_FMT

    for col_idx, col_letter in [(7, 'G'), (8, 'H'), (9, 'I'), (10, 'J')]:
        formula = f'=SUMPRODUCT((F{data_start_row}:F{data_end_row})*({col_letter}{data_start_row}:{col_letter}{data_end_row}))'
        ws.cell(row=summary_row, column=col_idx, value=formula)
        ws.cell(row=summary_row, column=col_idx).number_format = NUM_FMT

    apply_styles(ws, data_start_row, data_end_row, summary_row)

    merge_consecutive_cells(ws, 1, data_start_row, data_end_row)
    merge_consecutive_cells(ws, 2, data_start_row, data_end_row)

    return summary_row


def write_summary_sheet(wb, config, product_summaries):
    ws = wb.create_sheet(title="项目汇总")

    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value=f"{config['project_name']} - 硬件资源清单汇总")
    ws.cell(row=row, column=1).font = TITLE_FONT
    ws.cell(row=row, column=1).alignment = Alignment(horizontal='center', vertical='center')

    row = 3
    info_pairs = [
        ("项目名称", config['project_name']),
        ("建设范围", config.get('scope', '')),
        ("常住人口", f"{config.get('population', '')}万"),
        ("信创环境", "是" if config.get('xinchuang') else "否"),
    ]
    for label, value in info_pairs:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    for product_key, info in product_summaries.items():
        ws.cell(row=row, column=1, value=f"{product_key}档位").font = Font(bold=True)
        ws.cell(row=row, column=2, value=info.get('tier_desc', ''))
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value="各产品服务器资源汇总（数量×配置）")
    ws.cell(row=row, column=1).font = Font(bold=True, size=12)
    ws.cell(row=row, column=1).alignment = Alignment(horizontal='center', vertical='center')

    row += 1
    summary_headers = ["产品", "服务器数量", "CPU（核）", "内存（GB）", "系统盘（GB）", "数据盘（GB）"]
    for col_idx, header in enumerate(summary_headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center')

    first_data_row = row + 1

    for product_key, info in product_summaries.items():
        row += 1
        ws.cell(row=row, column=1, value=info['display_name']).border = THIN_BORDER

        sheet_name = info['sheet_name']
        summary_r = info['summary_row']

        refs = [
            (2, 'F'), (3, 'G'), (4, 'H'), (5, 'I'), (6, 'J')
        ]
        for col_idx, src_col in refs:
            ref = f"'{sheet_name}'!{src_col}{summary_r}"
            cell = ws.cell(row=row, column=col_idx, value=f"={ref}")
            cell.number_format = NUM_FMT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal='center')

    last_data_row = row

    if last_data_row >= first_data_row:
        row += 1
        ws.cell(row=row, column=1, value="合计").font = SUMMARY_FONT
        ws.cell(row=row, column=1).fill = SUMMARY_FILL
        ws.cell(row=row, column=1).border = THIN_BORDER
        ws.cell(row=row, column=1).alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=1)

        for col_idx in range(2, 7):
            col_letter = get_column_letter(col_idx)
            formula = f'=SUM({col_letter}{first_data_row}:{col_letter}{last_data_row})'
            cell = ws.cell(row=row, column=col_idx, value=formula)
            cell.number_format = NUM_FMT
            cell.font = SUMMARY_FONT
            cell.fill = SUMMARY_FILL
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal='center')

    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 14


def generate(config):
    TIER_FALLBACK_NOTES.clear()
    products = config.get('products', {})
    xinchuang = config.get('xinchuang', True)
    population = config.get('population')
    project_name = config.get('project_name', '未命名项目')
    scope = config.get('scope', '')

    wb = Workbook()
    wb.remove(wb.active)

    wxp_summary_row = write_product_sheet(
        wb, "WXP必选",
        f"{project_name} - WXP运维平台（必选）",
        "WXP为所有项目必选公共组件",
        load_wxp_rows(),
        is_wxp=True
    )

    selected_with_shared = [p for p in SHARED_COMPONENT_PRODUCTS if p in products]
    has_shared = len(selected_with_shared) >= 2

    product_summaries = {
        "WXP": {
            'sheet_name': "WXP必选",
            'summary_row': wxp_summary_row,
            'display_name': "WXP运维平台",
            'tier_desc': "必选公共组件",
        }
    }

    for product_key, params in products.items():
        if product_key == "区域平台":
            continue

        outpatient = params.get('outpatient')
        minimal = params.get('minimal', False)
        variant = params.get('variant')
        tier = determine_tier(product_key, outpatient=outpatient, population=population, minimal=minimal)

        raw_rows = load_product_data(product_key, xinchuang, tier, variant=variant)

        if has_shared and product_key != "智慧公卫":
            raw_rows = [r for r in raw_rows if not is_shared_component(r)]

        tier_desc = get_tier_description(product_key, tier, outpatient=outpatient, population=population)
        sheet_name = PRODUCT_SHEET_NAMES[product_key]

        if product_key == "基层医疗":
            if tier <= 1:
                full_title = f"{project_name} - 基层医疗信息系统（含基层卫生框架、电子病历）"
            else:
                full_title = f"{project_name} - 基层医疗信息系统"
        elif product_key == "智慧公卫":
            variant_str = f"（{variant}）" if variant else ""
            full_title = f"{project_name} - 智慧公卫系统{variant_str}"
        else:
            full_title = f"{project_name} - {sheet_name}"

        subtitle = f"档位：{tier_desc}"

        xinchuang_str = "信创" if xinchuang else "非信创"
        source_sheet = SHEET_NAME_TEMPLATES[product_key].format(xc=xinchuang_str)
        if source_sheet in TIER_FALLBACK_NOTES:
            orig_t, actual_t = TIER_FALLBACK_NOTES[source_sheet]
            subtitle += f" | ⚠️源数据无第{orig_t}档，数据回退自第{actual_t}档"

        summary_row = write_product_sheet(wb, sheet_name, full_title, subtitle, raw_rows)

        product_summaries[product_key] = {
            'sheet_name': sheet_name,
            'summary_row': summary_row,
            'display_name': sheet_name,
            'tier_desc': tier_desc,
        }

    if "区域平台" in products:
        combo = products["区域平台"].get('combo', '全量清单')
        tier = determine_tier("区域平台", population=population)
        raw_rows = load_regional_data(combo, xinchuang, tier)

        sheet_name = REGIONAL_SHEET_NAMES.get(combo, "区域平台")
        full_title = f"{project_name} - {sheet_name}"
        tier_desc = get_tier_description("区域平台", tier, population=population)
        subtitle = f"档位：{tier_desc}"

        summary_row = write_product_sheet(wb, sheet_name, full_title, subtitle, raw_rows)

        product_summaries["区域平台"] = {
            'sheet_name': sheet_name,
            'summary_row': summary_row,
            'display_name': sheet_name,
            'tier_desc': tier_desc,
        }

    write_summary_sheet(wb, config, product_summaries)

    return wb


def main():
    import argparse
    parser = argparse.ArgumentParser(description='WiNEX 硬件资源清单生成器')
    parser.add_argument('--config', required=True, help='配置文件路径 (JSON)')
    parser.add_argument('--source-dir', required=True,
                        help='平铺源数据 md 目录（xlsx_to_md.py 转换产物，如知识库清单笔记所在目录）')
    parser.add_argument('--output', help='输出文件路径 (xlsx)')
    args = parser.parse_args()

    with open(args.config, 'r', encoding='utf-8') as f:
        config = json.load(f)

    load_flat_source(args.source_dir)
    wb = generate(config)

    project_name = config.get('project_name', 'output')
    output_dir = Path.cwd() / project_name
    output_dir.mkdir(parents=True, exist_ok=True)

    if args.output:
        output_path = Path(args.output)
        if not output_path.is_absolute():
            output_path = output_dir / output_path
    else:
        output_path = output_dir / f"{project_name}_硬件资源清单.xlsx"

    wb.save(str(output_path))
    print(f"生成完成: {output_path}")


if __name__ == '__main__':
    main()
